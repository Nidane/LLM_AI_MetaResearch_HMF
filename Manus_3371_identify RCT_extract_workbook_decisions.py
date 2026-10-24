#!/usr/bin/env python3
"""Export retained workbook data to auditable CSV files.

The original screening workbook is a binary Excel file.  This program exports
*all four retained sheets* to editable/reviewable CSV source data and also
keeps the two concise decision-log CSVs used in the first reconstruction
package.  It makes no new eligibility decisions and does not modify the
reference workbook.

Run example
-----------
python src/extract_workbook_decisions.py \
  --workbook reference/HMF_RCT_extraction_for_NEC_sepsis.xlsx \
  --output-dir data
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence, Tuple

from openpyxl import load_workbook


WORKBOOK_SHEETS: Tuple[Tuple[str, str], ...] = (
    ("Eligible_RCTs", "eligible_rcts_full.csv"),
    ("Methods_and_notes", "methods_and_notes.csv"),
    ("Excluded_or_secondary", "excluded_or_secondary.csv"),
    ("Initial_keyword_candidates", "initial_keyword_candidates.csv"),
)


def as_text(value: Any) -> str:
    """Provide a stable representation of workbook cells for source CSV files."""
    return "" if value is None else str(value).replace("\r\n", "\n").replace("\r", "\n")


def sheet_data(workbook_path: Path, sheet: str) -> Tuple[List[str], List[Dict[str, str]]]:
    """Return the exact header and all nonblank rows from one workbook sheet."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet not in workbook.sheetnames:
        raise ValueError(f"Sheet '{sheet}' not present. Available: {workbook.sheetnames}")

    worksheet = workbook[sheet]
    header = [as_text(cell.value) for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    if any(not name for name in header):
        raise ValueError(f"Sheet '{sheet}' has a blank header; cannot safely export to CSV")

    rows: List[Dict[str, str]] = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        row = {column: as_text(value) for column, value in zip(header, values)}
        if any(row.values()):
            rows.append(row)
    return header, rows


def write_csv(path: Path, fieldnames: Sequence[str], rows: Iterable[Dict[str, str]]) -> int:
    """Write UTF-8 CSV using the provided source-data column order."""
    count = 0
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fieldnames), extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({name: row.get(name, "") for name in fieldnames})
            count += 1
    return count


def classify_decision(text: str) -> str:
    """Give a concise protocol-scope label without replacing source wording."""
    low = text.lower()
    if "borderline/exclude" in low:
        return "borderline; exclude from strict preterm/VLBW fortifier-only primary analysis"
    if "direct fortifier-source" in low or "direct hmbf vs bmbf" in low:
        return "include; direct fortifier-source RCT"
    if "ehmd" in low or "human milk diet" in low:
        return "include only if protocol accepts broader EHMD/bovine-product diet comparison"
    if low.startswith("include"):
        return "include; requires outcome-table/full-text verification for quantitative pooling"
    return "see original manual decision"


def write_concise_decision_logs(output_dir: Path, workbook_path: Path) -> Tuple[int, int]:
    """Write compact, convenient copies of the full retained decision sheets."""
    eligible_header, eligible_rows = sheet_data(workbook_path, "Eligible_RCTs")
    eligible_fields = [
        "screening_disposition",
        "inclusion_decision",
        "short_citation",
        "record_no",
        "pmid",
        "pmcid",
        "doi",
        "pubmed_title",
        "design",
        "population",
        "human_intervention",
        "comparator",
        "directness",
        "outcomes_in_abstract",
        "nec_human",
        "nec_bovine",
        "nec_definition",
        "sepsis_human",
        "sepsis_bovine",
        "sepsis_definition",
        "caveat",
    ]
    # Stop early if a future workbook drops one of these expected fields.
    missing = set(eligible_fields[1:]) - set(eligible_header)
    if missing:
        raise ValueError(f"Eligible_RCTs unexpectedly lacks: {sorted(missing)}")
    normalised_rows = []
    for row in eligible_rows:
        copied = dict(row)
        copied["screening_disposition"] = classify_decision(row.get("inclusion_decision", ""))
        normalised_rows.append(copied)
    n_eligible = write_csv(output_dir / "manual_eligible_or_borderline_decisions.csv", eligible_fields, normalised_rows)

    excluded_header, excluded_rows = sheet_data(workbook_path, "Excluded_or_secondary")
    excluded_fields = ["status", "reason", "identifier"]
    missing = set(excluded_fields) - set(excluded_header)
    if missing:
        raise ValueError(f"Excluded_or_secondary unexpectedly lacks: {sorted(missing)}")
    n_excluded = write_csv(output_dir / "manual_excluded_or_secondary_reports.csv", excluded_fields, excluded_rows)
    return n_eligible, n_excluded


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument(
        "--no-concise-logs",
        action="store_true",
        help="Export only full-sheet CSV source data, not the two compact decision logs.",
    )
    args = parser.parse_args()

    if not args.workbook.is_file():
        raise SystemExit(f"Workbook does not exist: {args.workbook}")
    args.output_dir.mkdir(parents=True, exist_ok=True)

    for sheet_name, filename in WORKBOOK_SHEETS:
        header, rows = sheet_data(args.workbook, sheet_name)
        count = write_csv(args.output_dir / filename, header, rows)
        print(f"Exported {count} rows from {sheet_name} -> {args.output_dir / filename}")

    if not args.no_concise_logs:
        n_eligible, n_excluded = write_concise_decision_logs(args.output_dir, args.workbook)
        print(f"Exported {n_eligible} eligible/borderline decision rows")
        print(f"Exported {n_excluded} excluded/secondary-report rows")


if __name__ == "__main__":
    main()
