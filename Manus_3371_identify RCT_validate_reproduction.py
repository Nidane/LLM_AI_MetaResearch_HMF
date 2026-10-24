#!/usr/bin/env python3
"""Validate the reconstructed HMF screening package and rebuilt workbook.

Checks performed:
- source-export hash matches the screening run;
- canonical parser produced exactly 3,371 supplied records;
- all retained eligible/borderline PMIDs appear in the broad candidate queue;
- all four source workbook sheets and both concise decision logs exist with the
  expected retained row counts;
- if supplied, rebuilt workbook data and sheet order match the retained workbook.

The test intentionally does not require byte-for-byte Excel identity because
metadata and ZIP ordering can legitimately differ between workbook builds.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook


CSV_EXPECTATIONS: Tuple[Tuple[str, int], ...] = (
    ("eligible_rcts_full.csv", 9),
    ("methods_and_notes.csv", 5),
    ("excluded_or_secondary.csv", 7),
    ("initial_keyword_candidates.csv", 53),
    ("manual_eligible_or_borderline_decisions.csv", 9),
    ("manual_excluded_or_secondary_reports.csv", 7),
    ("outcome_extraction_ledger.csv", 27),
    ("candidate_reconciliation.csv", 53),
)
EXPECTED_ELIGIBLE_OR_BORDERLINE_PMIDS = {
    "39660504", "42045666", "38545091", "36857051", "36528055",
    "35809401", "29878061", "23968744", "20036378",
}


def checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalise_cell(value: Any) -> str:
    """Compare workbook values semantically after CSV's string serialization."""
    return "" if value is None else str(value).replace("\r\n", "\n").replace("\r", "\n")


def values_from_sheet(path: Path, sheet: str) -> List[Tuple[str, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    return [tuple(normalise_cell(cell) for cell in row) for row in workbook[sheet].iter_rows(values_only=True)]


def row_count(csv_path: Path) -> int:
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        return sum(1 for _ in csv.DictReader(handle))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--data-dir", required=True, type=Path)
    parser.add_argument("--reference-workbook", required=True, type=Path)
    parser.add_argument("--rebuilt-workbook", required=False, type=Path)
    args = parser.parse_args()

    errors: List[str] = []
    if not args.input.is_file():
        errors.append(f"Input missing: {args.input}")
    if not args.data_dir.is_dir():
        errors.append(f"Data directory missing: {args.data_dir}")
    if not args.reference_workbook.is_file():
        errors.append(f"Reference workbook missing: {args.reference_workbook}")
    if errors:
        raise SystemExit("\n".join(errors))

    summary_path = args.data_dir / "screening_summary.json"
    if not summary_path.is_file():
        errors.append(f"Missing automated-screen metadata: {summary_path}")
    else:
        summary: Dict[str, Any] = json.loads(summary_path.read_text(encoding="utf-8"))
        if summary.get("input_sha256") != checksum(args.input):
            errors.append("Input SHA-256 does not match screening_summary.json")
        parsing = summary.get("parsing", {})
        if summary.get("parse_mode") != "canonical":
            errors.append("Screen did not use canonical parse mode")
        if int(parsing.get("canonical_records", 0)) != 3371:
            errors.append("Unexpected canonical record count (expected 3371 for retained input)")
        if int(summary.get("broad_manual_candidate_count", 0)) < 1:
            errors.append("Broad manual candidate count is zero")

    candidate_path = args.data_dir / "rct_candidate_flags.tsv"
    if not candidate_path.is_file():
        errors.append(f"Missing candidate table: {candidate_path}")
    else:
        with candidate_path.open(encoding="utf-8", newline="") as handle:
            candidates = list(csv.DictReader(handle, delimiter="\t"))
        candidate_pmids = {row.get("pmid", "") for row in candidates}
        absent = EXPECTED_ELIGIBLE_OR_BORDERLINE_PMIDS - candidate_pmids
        if absent:
            errors.append(f"Final-workbook PMID(s) absent from candidate TSV: {sorted(absent)}")

    for name, expected_rows in CSV_EXPECTATIONS:
        path = args.data_dir / name
        if not path.is_file():
            errors.append(f"Missing expected CSV: {name}")
        else:
            actual_rows = row_count(path)
            if actual_rows != expected_rows:
                errors.append(f"Unexpected row count in {name}: expected {expected_rows}, got {actual_rows}")

    if args.rebuilt_workbook:
        if not args.rebuilt_workbook.is_file():
            errors.append(f"Rebuilt workbook missing: {args.rebuilt_workbook}")
        else:
            original = load_workbook(args.reference_workbook, read_only=True, data_only=True)
            rebuilt = load_workbook(args.rebuilt_workbook, read_only=True, data_only=True)
            if original.sheetnames != rebuilt.sheetnames:
                errors.append(f"Workbook sheet order/names differ: {original.sheetnames} vs {rebuilt.sheetnames}")
            else:
                for sheet in original.sheetnames:
                    if values_from_sheet(args.reference_workbook, sheet) != values_from_sheet(args.rebuilt_workbook, sheet):
                        errors.append(f"Rebuilt workbook values differ in sheet: {sheet}")

    if errors:
        raise SystemExit("VALIDATION FAILED:\n- " + "\n- ".join(errors))
    print("VALIDATION PASS")
    print("Input hash matches; canonical parser recovered the 3,371 numbered records.")
    print("Candidate list contains all 9 retained eligible/borderline PMIDs.")
    print("All four source workbook sheets and both concise decision logs have expected row counts.")
    if args.rebuilt_workbook:
        print("Rebuilt workbook matches retained workbook values and sheet order.")


if __name__ == "__main__":
    main()
