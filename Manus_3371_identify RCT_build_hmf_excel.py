#!/usr/bin/env python3
"""Build an editable HMF RCT screening workbook from the exported CSV source data.

This script reproduces the retained workbook's data structure (four sheets,
headers, column widths, frozen header row, wrap text, and header style) from
plain CSV source files.  It does not claim byte-for-byte identity with the
retained original .xlsx; metadata and other non-substantive workbook details
can differ.

Run example
-----------
python src/build_hmf_excel.py \
  --data-dir output_example \
  --output HMF_RCT_extraction_rebuilt.xlsx
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# Kept in the source code because these are presentation settings—not evidence
# data—and they match the retained workbook's visible structure.
SHEET_SPECS: Tuple[Dict[str, object], ...] = (
    {
        "sheet": "Eligible_RCTs",
        "csv": "eligible_rcts_full.csv",
        "widths": [
            55, 33, 12, 12, 13, 35, 55, 55, 41, 55, 55, 12, 12, 12,
            55, 55, 55, 55, 12, 12, 55, 14, 15, 55, 17, 18, 55, 55, 55,
        ],
    },
    {
        "sheet": "Methods_and_notes",
        "csv": "methods_and_notes.csv",
        "widths": [25, 55],
    },
    {
        "sheet": "Excluded_or_secondary",
        "csv": "excluded_or_secondary.csv",
        "widths": [55, 55, 55],
    },
    {
        "sheet": "Initial_keyword_candidates",
        "csv": "initial_keyword_candidates.csv",
        "widths": [12, 12, 35, 55, 55, 55],
    },
)

HEADER_FILL = PatternFill(fill_type="solid", fgColor="4F81BD")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(wrap_text=True, vertical="top")
CELL_ALIGNMENT = Alignment(wrap_text=True, vertical="top")


def read_csv(path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    """Read one UTF-8 source CSV and validate its header is present."""
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError(f"No header row found in {path}")
        header = list(reader.fieldnames)
        return header, list(reader)


def build_sheet(workbook: Workbook, spec: Dict[str, object], data_dir: Path) -> int:
    """Build a single sheet using its source CSV and retained presentation specs."""
    sheet_name = str(spec["sheet"])
    input_csv = data_dir / str(spec["csv"])
    if not input_csv.is_file():
        raise FileNotFoundError(f"Required source data file does not exist: {input_csv}")

    header, rows = read_csv(input_csv)
    widths = list(spec["widths"])
    if len(header) != len(widths):
        raise ValueError(
            f"{input_csv} has {len(header)} header columns but {sheet_name} requires {len(widths)} widths"
        )

    worksheet = workbook.create_sheet(title=sheet_name)
    worksheet.append(header)
    for row in rows:
        worksheet.append([row.get(column, "") for column in header])

    worksheet.freeze_panes = "A2"
    for idx, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(idx)].width = width
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = CELL_ALIGNMENT
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", required=True, type=Path, help="Folder produced by extract_workbook_decisions.py")
    parser.add_argument("--output", required=True, type=Path, help="Destination .xlsx path")
    args = parser.parse_args()

    if not args.data_dir.is_dir():
        raise SystemExit(f"Data directory does not exist: {args.data_dir}")
    args.output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)
    total = 0
    for spec in SHEET_SPECS:
        count = build_sheet(workbook, spec, args.data_dir)
        total += count
        print(f"Built {spec['sheet']}: {count} data rows")
    workbook.save(args.output)
    print(f"Saved {args.output} ({total} data rows across {len(SHEET_SPECS)} sheets)")


if __name__ == "__main__":
    main()
